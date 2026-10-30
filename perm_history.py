"""Official DOL PERM employer history: exact-name matching only.

Refresh with: python perm_history.py refresh
No H-1B inference; case determinations do not guarantee future sponsorship.
"""
import argparse
import gzip
import json
import re
import tempfile
from collections import defaultdict
from datetime import datetime, date, timezone
from pathlib import Path
from urllib.parse import urlparse
import requests
from openpyxl import load_workbook

INDEX=Path("perm_history_index.json.gz")
OFFICIAL="https://www.dol.gov/agencies/eta/foreign-labor/performance"
SOURCES={
 "FY2024-old":"https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/PERM_Disclosure_Data_FY2024_Q4.xlsx",
 "FY2024-new":"https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/PERM_Disclosure_Data_New_Form_FY2024_Q4.xlsx",
 "FY2025":"https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/PERM_Disclosure_Data_FY2025_Q4.xlsx",
 "FY2026-Q3":"https://www.dol.gov/media/PERM_Disclosure_Data_FY2026_Q3.xlsx",
}
LEGAL={"INC","INCORPORATED","LLC","L L C","LTD","LIMITED","CORP","CORPORATION","CO","COMPANY","LP","LLP","THE"}
def normalize(name):
    parts=re.findall(r"[A-Z0-9]+",str(name or "").upper().replace("&"," AND "))
    while parts and parts[-1] in LEGAL:
        parts.pop()
    if parts and parts[0]=="THE":
        parts.pop(0)
    return " ".join(parts)

def column(headers,*choices):
    mapping={re.sub(r"[^A-Z0-9]","",str(v or "").upper()):i for i,v in enumerate(headers)}
    for choice in choices:
        key=re.sub(r"[^A-Z0-9]","",choice.upper())
        if key in mapping:return mapping[key]
    return None

def parse_date(value):
    if isinstance(value,(datetime,date)):
        return value.isoformat()[:10]
    if isinstance(value,str) and value.strip():
        x=value.strip()
        for fmt in ("%Y-%m-%d","%m/%d/%Y","%Y-%m-%d %H:%M:%S"):
            try:return datetime.strptime(x,fmt).strftime("%Y-%m-%d")
            except ValueError:pass
    return ""

def records_from_workbook(path):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        for sheet in book.worksheets:
            iterator=sheet.iter_rows(values_only=True)
            headers=next(iterator,None)
            if not headers:continue
            employer=column(headers,"EMPLOYER_NAME","EMPLOYER_NAME_1")
            status=column(headers,"CASE_STATUS","CASE_STATUS_1")
            case=column(headers,"CASE_NUMBER","CASE_NO")
            decided=column(headers,"DECISION_DATE","DECISION_DATE_1")
            if employer is None or status is None:
                continue
            for row in iterator:
                if not row:continue
                name=str(row[employer] or "").strip()
                if not name:continue
                yield (name,str(row[status] or "").strip().upper(),
                       str(row[case] or "").strip() if case is not None else "",
                       parse_date(row[decided]) if decided is not None else "")
    finally:
        book.close()

def aggregate_batches(batches):
    """Input: (source name, iterable of (employer,status,case,date))."""
    summary={}
    seen=set()
    for source,rows in batches:
        year=re.search(r"FY(\d{4})",source)
        if not year:continue
        fiscal=year.group(1)
        for name,status,case,determined in rows:
            key=normalize(name)
            if not key:continue
            # DOL FY2024 has legacy and new-form files; avoid duplicate cases.
            identity=(fiscal,case) if case else (fiscal,key,status,determined)
            if identity in seen:continue
            seen.add(identity)
            record=summary.setdefault(key,{"display_name":name,"determinations":0,
                "certified":0,"by_year":{},"latest_determination":"","variants":[]})
            record["determinations"]+=1
            if status=="CERTIFIED":record["certified"]+=1
            record["by_year"].setdefault(fiscal,{"determinations":0,"certified":0})
            record["by_year"][fiscal]["determinations"]+=1
            if status=="CERTIFIED":record["by_year"][fiscal]["certified"]+=1
            if determined>record["latest_determination"]:
                record["latest_determination"]=determined
            if name not in record["variants"] and len(record["variants"])<5:
                record["variants"].append(name)
    return summary

def refresh(index=INDEX,session=None):
    session=session or requests.Session()
    batches=[]
    paths=[]
    try:
        with tempfile.TemporaryDirectory() as temp:
            for key,url in SOURCES.items():
                if urlparse(url).hostname!="www.dol.gov":
                    raise ValueError("Non-DOL source disallowed")
                path=Path(temp)/(key+".xlsx")
                with session.get(url,stream=True,timeout=(15,100),headers={"User-Agent":"Mozilla/5.0"}) as response:
                    response.raise_for_status()
                    size=0
                    with path.open("wb") as out:
                        for chunk in response.iter_content(1024*1024):
                            size+=len(chunk)
                            if size>150_000_000:raise ValueError("DOL workbook exceeds safety size limit")
                            out.write(chunk)
                paths.append((key,path))
            batches=[(key,records_from_workbook(path)) for key,path in paths]
            employers=aggregate_batches(batches)
            if not employers:raise ValueError("DOL workbooks produced no usable PERM records")
            payload={"generated_utc":datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "source_page":OFFICIAL,"source_files":SOURCES,
                "years":[2024,2025,2026],"employers":employers}
            target=Path(index)
            temp_index=target.with_suffix(".tmp")
            with gzip.open(temp_index,"wt",encoding="utf-8") as out:
                json.dump(payload,out,ensure_ascii=False,separators=(",",":"))
            temp_index.replace(target)
            print(f"Indexed {len(employers)} distinct PERM employer names from official DOL disclosures")
            return payload
    finally:
        if hasattr(session,"close") and session is not requests:
            session.close()

def load(index=INDEX):
    path=Path(index)
    if not path.exists():return None
    try:
        with gzip.open(path,"rt",encoding="utf-8") as stream:
            payload=json.load(stream)
        return payload if isinstance(payload.get("employers"),dict) else None
    except (OSError,ValueError,TypeError):return None

def lookup(company,index):
    """Exact normalized employer name ONLY; no invented parent/subsidiary match."""
    result={"perm_status":"Unknown","perm_certified_3yr":0,"perm_cases_3yr":0,
            "perm_latest_determination":"","perm_matched_employer":"",
            "perm_source":OFFICIAL,"perm_data_through":"FY2026 Q3"}
    if not index:return result
    employer=index.get("employers",{}).get(normalize(company))
    if not employer:return result
    result.update({"perm_status":"Recent PERM certifications" if employer["certified"] else "PERM determinations found (no certifications)",
        "perm_certified_3yr":employer["certified"],
        "perm_cases_3yr":employer["determinations"],
        "perm_latest_determination":employer.get("latest_determination",""),
        "perm_matched_employer":employer.get("display_name","")})
    return result

if __name__=="__main__":
    args=argparse.ArgumentParser()
    args.add_argument("action",choices=["refresh"])
    parsed=args.parse_args()
    if parsed.action=="refresh":refresh()
